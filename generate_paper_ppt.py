import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette
    COLOR_BG_DARK = RGBColor(15, 32, 39)       # #0F2027
    COLOR_PRIMARY = RGBColor(24, 144, 255)     # Tech Blue #1890FF
    COLOR_SECONDARY = RGBColor(44, 83, 100)    # Slate Blue
    COLOR_TEXT_DARK = RGBColor(38, 38, 38)     # #262626
    COLOR_TEXT_MUTED = RGBColor(115, 115, 115) # Gray
    COLOR_TEXT_LIGHT = RGBColor(255, 255, 255) # White
    COLOR_CARD_BG = RGBColor(247, 249, 252)    # Light gray-blue
    COLOR_GREEN = RGBColor(56, 158, 13)        # #389E0D
    COLOR_RED = RGBColor(207, 19, 34)          # #CF1322
    COLOR_GOLD = RGBColor(212, 136, 6)         # Amber

    def add_header(slide, title_text, category_text="MASK: RESEARCH PRESENTATION"):
        # Header banner shape
        banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.15))
        banner.fill.solid()
        banner.fill.fore_color.rgb = COLOR_BG_DARK
        banner.line.fill.background()

        # Category
        tb_cat = slide.shapes.add_textbox(Inches(0.8), Inches(0.12), Inches(11), Inches(0.3))
        tf_cat = tb_cat.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_PRIMARY
        p_cat.font.name = "Microsoft YaHei"

        # Title
        tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.38), Inches(11.5), Inches(0.6))
        tf_title = tb_title.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_TEXT_LIGHT
        p_title.font.name = "Microsoft YaHei"

    def add_card(slide, left, top, width, height, bg_color=COLOR_CARD_BG, border_color=None):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1.5)
        else:
            shape.line.fill.background()
        return shape

    # ==========================================
    # SLIDE 1: 封面 Title Slide
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = COLOR_BG_DARK
    bg1.line.fill.background()

    # Title Card
    tb = s1.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(11.333), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "MASK: 面向不完全信息博弈的\n风险门控动作层塑形决策框架"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_LIGHT
    p.font.name = "Microsoft YaHei"
    p.space_after = Pt(14)

    p2 = tf.add_paragraph()
    p2.text = "Risk-Gated Action-Level Shaping for Imperfect-Information Games"
    p2.font.size = Pt(18)
    p2.font.color.rgb = COLOR_PRIMARY
    p2.font.name = "Calibri"
    p2.space_after = Pt(28)

    p3 = tf.add_paragraph()
    p3.text = "• 创新提出：三路动作解耦 (Exploit / Safe / Deceive) + 神经门控长远估值 + 边际护栏仲裁\n• 严谨实证：320,000 局正式重跑、20 个独立 Outer Seed、CRN 全配对审计\n• 核心成果：平均分提升 +5.565，放铳率下降 -1.545pp，诱导荣和率暴增 +8.450pp，CVaR 提升 +15.685"
    p3.font.size = Pt(14)
    p3.font.color.rgb = RGBColor(200, 215, 225)
    p3.font.name = "Microsoft YaHei"

    # ==========================================
    # SLIDE 2: 研究背景与挑战
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "研究背景与核心挑战：为什么不完全信息博弈极具挑战？", "1. RESEARCH BACKGROUND")

    # 3 Cards
    cards_data = [
        ("极端状态与非对称信息", 
         "麻将具有超 10^48 级庞大状态空间，存在三家完全隐蔽的私有手牌与随机牌墙，信息极度非对称，无法像围棋等完全信息博弈进行全树前向搜寻。",
         RGBColor(24, 144, 255)),
        ("多方博弈与负和/非零和机制", 
         "4 人参与且相互制衡。点炮（Deal-in）直接导致自身重大失分并使特定对手获利，单纯追求牌效的策略极易沦为其他对手的'提款机'。",
         RGBColor(250, 140, 22)),
        ("大模型直接落地的三大瓶颈", 
         "1. 牌效幻觉：过分执着自身向听数，忽视隐蔽风险\n2. 盲目欺骗：无门控 Bluff 导致放铳率激增，CVaR 崩溃\n3. 延迟过高：全向 MCTS 难以满足毫秒级准实时决策",
         RGBColor(245, 34, 45)),
    ]

    for i, (title, content, bar_color) in enumerate(cards_data):
        left = Inches(0.8 + i * 4.0)
        card = add_card(s2, left, Inches(1.6), Inches(3.73), Inches(5.2))
        
        # Color top bar
        bar = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, Inches(1.6), Inches(3.73), Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = bar_color
        bar.line.fill.background()

        tb = s2.shapes.add_textbox(left + Inches(0.2), Inches(1.9), Inches(3.33), Inches(4.7))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXT_DARK
        p.font.name = "Microsoft YaHei"
        p.space_after = Pt(14)

        p2 = tf.add_paragraph()
        p2.text = content
        p2.font.size = Pt(13)
        p2.font.color.rgb = COLOR_TEXT_MUTED
        p2.font.name = "Microsoft YaHei"

    # ==========================================
    # SLIDE 3: 核心动机与关键洞见
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "核心动机与博弈洞见：欺骗的本质在于'诱捕'而非'驱赶'", "2. MOTIVATION & INSIGHT")

    # Left: The False Belief in Poker vs Mahjong
    add_card(s3, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), border_color=COLOR_RED)
    tb = s3.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.2), Inches(4.8))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "传统偏见：欺骗是为了迫使对手放弃 (Fold)？"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = COLOR_RED
    p.space_after = Pt(12)

    bullets_left = [
        "德州扑克逻辑：通过强力下注（Bluff）迫使对手弃牌即可赢走底池。",
        "麻将博弈悖论：如果智能体把所有对手都吓得严密防守，对手将只出绝对安全牌，导致牌局走向流局（荒牌）。",
        "实测惨痛教训：消融实验中无脑常开欺骗（Constant-deceive）取得了最高的严格 FFR (40.02%)，但平均得分从 24.89 分暴跌至 13.79 分（全线崩盘）！"
    ]
    for b in bullets_left:
        p_b = tf.add_paragraph()
        p_b.text = "• " + b
        p_b.font.size = Pt(13)
        p_b.font.color.rgb = COLOR_TEXT_DARK
        p_b.space_after = Pt(10)

    # Right: MASK Insight
    add_card(s3, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), border_color=COLOR_GREEN)
    tb_r = s3.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.3), Inches(4.8))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True
    p_r = tf_r.paragraphs[0]
    p_r.text = "MASK 核心洞见：动作塑形旨在'诱导放铳' (Ron)"
    p_r.font.size = Pt(17)
    p_r.font.bold = True
    p_r.font.color.rgb = COLOR_GREEN
    p_r.space_after = Pt(12)

    bullets_right = [
        "高阶诱导策略：在自身听牌或逼近听牌时，打出传递'假安全信号'的牌，诱骗对手误以为安全而继续进攻打出我方目标牌。",
        "精准门控裁决：绝不为了欺骗而破坏自身牌效；只有门控网络判断欺骗带来的边际收益明确大于进攻时才激活。",
        "实测显著战果：MASK-full 相比无欺骗基线，诱导荣和率 (excess-DIR) 暴增 +8.450pp (p < 1e-14)，净胜分提升 +5.565 分！"
    ]
    for b in bullets_right:
        p_b = tf_r.add_paragraph()
        p_b.text = "✔ " + b
        p_b.font.size = Pt(13)
        p_b.font.color.rgb = COLOR_TEXT_DARK
        p_b.space_after = Pt(10)

    # ==========================================
    # SLIDE 4: 相关工作对比表
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "相关工作对比：不同技术路径的特性与局限分析", "3. RELATED WORK")

    rows = 5
    cols = 5
    table_shape = s4.shapes.add_table(rows, cols, Inches(0.8), Inches(1.6), Inches(11.733), Inches(5.1))
    t = table_shape.table

    t.columns[0].width = Inches(2.3)
    t.columns[1].width = Inches(2.3)
    t.columns[2].width = Inches(2.5)
    t.columns[3].width = Inches(2.3)
    t.columns[4].width = Inches(2.333)

    headers = ["方法类别", "代表系统", "核心机制", "核心优势", "关键局限 / 痛点"]
    for j, h in enumerate(headers):
        cell = t.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_BG_DARK
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_TEXT_LIGHT
        p.alignment = PP_ALIGN.CENTER

    data = [
        ("深度强化学习 (DRL)", "Suphx, Mortal", "价值网络 + 自博弈搜索", "牌效高，自博弈统治力强", "无显式信念建模，黑盒不可解，无法主动诱导对手"),
        ("子博弈求解 (PBS)", "ReBeL, DeepStack", "递归求解公共信念状态", "可证明的局部 Nash 保证", "状态空间随人数爆炸，麻将中推理延迟过高无法落地"),
        ("提示大模型 (LLM)", "Prompt GPT-4", "基于上下文的提示词工程", "具备规则理解与规划常识", "无法量化隐蔽风险，极易放铳点炮；无动作层概率塑形"),
        ("MASK (本文方法)", "MASK (Ours)", "LLM课程 + 神经门控 + 动作塑形", "兼顾高牌效推进与动态风控，诱导荣和率大增", "单步延迟低 (p50: 246ms)，兼顾安全与超额收益")
    ]

    for i, row in enumerate(data):
        for j, val in enumerate(row):
            cell = t.cell(i+1, j)
            cell.fill.solid()
            if i == 3: # Highlight MASK
                cell.fill.fore_color.rgb = RGBColor(230, 247, 255)
            else:
                cell.fill.fore_color.rgb = COLOR_CARD_BG if i % 2 == 0 else RGBColor(255, 255, 255)
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(11)
            p.font.name = "Microsoft YaHei"
            if i == 3:
                p.font.bold = True
                p.font.color.rgb = COLOR_PRIMARY if j == 0 else COLOR_TEXT_DARK
            else:
                p.font.color.rgb = COLOR_TEXT_DARK

    # ==========================================
    # SLIDE 5: 问题形式化与数学建模
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "问题形式化与数学建模：部分可观测随机博弈 (POSG)", "4. PROBLEM FORMULATION")

    # Math Card Left
    add_card(s5, Inches(0.8), Inches(1.6), Inches(4.5), Inches(5.2))
    tb_m = s5.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(4.1), Inches(4.8))
    tf_m = tb_m.text_frame
    tf_m.word_wrap = True
    p = tf_m.paragraphs[0]
    p.text = "形式化建模定义"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(12)

    defs = [
        "博弈元组：M = <N, S, {Ai}, P, {Ri}, {Oi}, γ>",
        "参与者集合：N = {0, 1, 2, 3}，其中 i=0 为我方智能体 P0，j ∈ {1,2,3} 为对手。",
        "局部观测：o_0,t = <h_0,t, I_t, M_0,t>，包含自身手牌 h_0,t 与全局公共弃牌河历史 I_t。",
        "公开信念分布：B_phi(I_t) 输出对手听牌概率 q_j,t = P(Shanten_j ≤ 0 | I_t)。",
        "优化目标：在最小化放铳风险与尾部损失 (CVaR) 前提下，最大化累积期望回报。"
    ]
    for d in defs:
        p_d = tf_m.add_paragraph()
        p_d.text = "• " + d
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = COLOR_TEXT_DARK
        p_d.space_after = Pt(10)

    # Table Right: Notation
    t_shape = s5.shapes.add_table(6, 3, Inches(5.5), Inches(1.6), Inches(7.0), Inches(5.2))
    t = t_shape.table
    t.columns[0].width = Inches(1.6)
    t.columns[1].width = Inches(1.6)
    t.columns[2].width = Inches(3.8)

    notations = [
        ("符号", "类型", "物理 / 数学含义"),
        ("h_0,t", "序列 / 离散", "P0 私有手牌集合及向听数 shanten(h_0,t)"),
        ("I_t", "序列集合", "公共弃牌河、各家副露、打牌时延与牌墙余量"),
        ("B_phi(I_t)", "概率向量", "公开信念估计器输出的对手听牌概率 q_j"),
        ("V_theta(s, a_m)", "实数标量", "门控网络对模式动作 a_m 预测的期望价值 Q"),
        ("excess-DIR", "百分比 (pp)", "相对无欺骗基准的超额诱导荣和率 (净胜分关键)")
    ]
    for i, (col1, col2, col3) in enumerate(notations):
        for j, val in enumerate([col1, col2, col3]):
            cell = t.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_BG_DARK if i == 0 else (COLOR_CARD_BG if i % 2 == 1 else RGBColor(255, 255, 255))
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(11)
            p.font.bold = (i == 0)
            p.font.color.rgb = COLOR_TEXT_LIGHT if i == 0 else COLOR_TEXT_DARK

    # ==========================================
    # SLIDE 6: MASK 整体系统架构
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "MASK 整体系统架构：感知、提案与门控三层协同体系", "5. OVERALL ARCHITECTURE")

    # 3 Large Horizontal Boxes
    box_data = [
        ("层级一：感知与隐蔽信念推断层 (Perception & Inference)",
         "• 输入源：私有手牌 + 公共牌河历史 + 对手副露与打牌节奏\n• 核心组件：B_phi 公开信念估计器 (AUC=0.7897) 提取对手听牌威胁 q_j\n• 启发特征：z_state 提取对手公共特征与牌河漂移度，融合为特征向量 X_t",
         RGBColor(24, 144, 255)),
        ("层级二：三路差异化动作提案层 (Multi-Mode Proposals)",
         "• ① Exploit 动作：LLM (Qwen-base 课程微调) 决策，以牌效进张与做大牌为目标\n• ② Safe 动作：防御算法根据 B_phi 与现物/筋牌库挑选点炮概率最低的安全张\n• ③ Deceive 动作：动作层塑形模块挑选向牌河释放'虚假安全感'的诱导弃牌",
         RGBColor(250, 140, 22)),
        ("层级三：神经门控网络与边际护栏层 (Arbitration & Guardrails)",
         "• GateValueNet：轻量双层 MLP 评估 V(exploit), V(safe), V(deceive) 期望价值\n• 边际阈值仲裁：严格要求 V(deceive) - V(exploit) ≥ τ_deceive，否则自动回退\n• 最终动作执行：毫秒级响应提交动作至游戏引擎，杜绝盲目欺骗与无脑防御",
         RGBColor(56, 158, 13))
    ]

    for i, (title, content, bar_color) in enumerate(box_data):
        top = Inches(1.6 + i * 1.8)
        add_card(s6, Inches(0.8), top, Inches(11.733), Inches(1.55))
        
        # Left color bar
        bar = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), top, Inches(0.15), Inches(1.55))
        bar.fill.solid()
        bar.fill.fore_color.rgb = bar_color
        bar.line.fill.background()

        tb = s6.shapes.add_textbox(Inches(1.2), top + Inches(0.1), Inches(11.1), Inches(1.35))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = bar_color
        p.font.name = "Microsoft YaHei"
        p.space_after = Pt(6)

        p2 = tf.add_paragraph()
        p2.text = content
        p2.font.size = Pt(12)
        p2.font.color.rgb = COLOR_TEXT_DARK
        p2.font.name = "Microsoft YaHei"

    # ==========================================
    # SLIDE 7: 核心方法：公开信念估计与动作提案
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "核心方法详解：公开信念估计 (B_phi) 与三路动作生成", "6. METHOD: INFERENCE & PROPOSALS")

    # Left: Belief Oracle
    add_card(s7, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    tb_l = s7.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.2), Inches(4.8))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True
    p = tf_l.paragraphs[0]
    p.text = "公开信息信念估计器 (B_phi)"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(10)

    b_pts = [
        "信息绝对合规：不访问对手私有手牌，仅使用公开牌河、副露、巡目与弃牌序列。",
        "网络模型架构：GBDT / MLP 配合 Isotonic 等保回归校准，极大消除极端自信过拟合。",
        "高精度预测验证 (E1)：Balanced 评测集 AUC 达到 0.7897，相对基础先验 Brier 分数降幅达 42.87% (p=5.45e-18)。",
        "角色定位：消融实验证实其不作为直接提分动力，而是提供最关键的'放铳安全侧防护' (DR 显著压低)。"
    ]
    for pt in b_pts:
        p_pt = tf_l.add_paragraph()
        p_pt.text = "• " + pt
        p_pt.font.size = Pt(12)
        p_pt.font.color.rgb = COLOR_TEXT_DARK
        p_pt.space_after = Pt(8)

    # Right: Multi-mode proposal
    add_card(s7, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    tb_r = s7.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.3), Inches(4.8))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True
    p = tf_r.paragraphs[0]
    p.text = "三路候选动作提案机制"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = COLOR_GOLD
    p.space_after = Pt(10)

    m_pts = [
        "① Exploit 提案 (牌效优先)：大模型根据当前手牌结构提出最优进张动作；若判定可胡牌 ('h') 则直接执行终结对局。",
        "② Safe 提案 (避险防守)：遍历手牌，结合 B_phi 风险评估和现物/筋牌/绝张规则，输出点炮危险度最小的牌。",
        "③ Deceive 提案 (诱导塑形)：在向听数 ≤ 1 且自身不至于严重损张的前提下，挑选向牌河释放迷惑信号的弃牌。",
        "独立候选解耦：三路动作并行提出，若欺骗动作与进攻动作相同或无效，欺骗通道自动关闭，降低决策干扰。"
    ]
    for pt in m_pts:
        p_pt = tf_r.add_paragraph()
        p_pt.text = "✔ " + pt
        p_pt.font.size = Pt(12)
        p_pt.font.color.rgb = COLOR_TEXT_DARK
        p_pt.space_after = Pt(8)

    # ==========================================
    # SLIDE 8: 核心方法：门控价值网络与边际护栏
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "核心方法详解：神经门控网络与边际安全护栏", "7. METHOD: GATE-VALUE & GUARDRAILS")

    # Top: GateValue Net
    add_card(s8, Inches(0.8), Inches(1.6), Inches(11.733), Inches(2.4))
    tb_t = s8.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(11.3), Inches(2.1))
    tf_t = tb_t.text_frame
    tf_t.word_wrap = True
    p = tf_t.paragraphs[0]
    p.text = "GateValueNet: 轻量高效的长远期望价值网络"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(6)

    p_body = tf_t.add_paragraph()
    p_body.text = (
        "• 输入向量融合：状态向量（手牌向听数、牌墙余量、风位）+ B_phi 对手听牌风险 + 三路动作的离散属性特征。\n"
        "• 结构极度轻量：采用 Linear(D, 128) -> ReLU -> Linear(128, 64) -> ReLU -> Linear(64, 3)，端到端推理毫秒级响应。\n"
        "• 输出目标对齐：分别输出当前局势下执行 Exploit, Safe, Deceive 对应的期望长期折现收益 V(m)。"
    )
    p_body.font.size = Pt(13)
    p_body.font.color.rgb = COLOR_TEXT_DARK

    # Bottom: Margin Guardrails
    add_card(s8, Inches(0.8), Inches(4.3), Inches(11.733), Inches(2.6))
    tb_b = s8.shapes.add_textbox(Inches(1.0), Inches(4.4), Inches(11.3), Inches(2.3))
    tf_b = tb_b.text_frame
    tf_b.word_wrap = True
    p = tf_b.paragraphs[0]
    p.text = "边际阈值保护与自适应回退机制 (Margin Guardrails)"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = COLOR_GREEN
    p.space_after = Pt(6)

    p_bbody = tf_b.add_paragraph()
    p_bbody.text = (
        "• 欺骗触发硬门槛：必须满足 V(deceive) - V(exploit) ≥ τ_deceive (默认 0.05)。唯有欺骗能带来显著超额收益时才允许激活！\n"
        "• 防守避险约束：必须满足 V(safe) - V(exploit) ≥ τ_safe 且外部威胁指数 ρ_t ≥ 阈值，杜绝过度防守导致弃和。\n"
        "• 优势不足自动回退：当网络判定欺骗收益微弱或风险模糊时，强制 Fallback 回退至 Exploit（LLM 进攻），确保系统稳健底线。"
    )
    p_bbody.font.size = Pt(13)
    p_bbody.font.color.rgb = COLOR_TEXT_DARK

    # ==========================================
    # SLIDE 9: 实验设置与大规模双盲审计
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "实验设置与大规模双盲审计：高标准科研实证规范", "8. EXPERIMENTAL PROTOCOL")

    # 4 Stat Metric Cards
    stats_data = [
        ("320,000 局", "正式重跑局数 (P4)\n零未结算对局 (unsettled=0)", COLOR_PRIMARY),
        ("20 个", "独立 Outer Seeds\n(2026082001-2026082020)", COLOR_GREEN),
        ("100% 配对", "Common Random Numbers\n跨系统逐局 Seed 严格对齐", COLOR_GOLD),
        ("零失败审计", "Artifact 哈希全程冻结\n严格 FFR 与终局核验全过", COLOR_RED),
    ]
    for i, (val, lbl, c) in enumerate(stats_data):
        left = Inches(0.8 + i * 3.0)
        card = add_card(s9, left, Inches(1.6), Inches(2.73), Inches(1.6))
        tb = s9.shapes.add_textbox(left, Inches(1.7), Inches(2.73), Inches(1.4))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = val
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = c
        p.alignment = PP_ALIGN.CENTER
        
        p2 = tf.add_paragraph()
        p2.text = lbl
        p2.font.size = Pt(11)
        p2.font.color.rgb = COLOR_TEXT_MUTED
        p2.alignment = PP_ALIGN.CENTER

    # Lower: Protocol Details
    add_card(s9, Inches(0.8), Inches(3.5), Inches(11.733), Inches(3.3))
    tb_proto = s9.shapes.add_textbox(Inches(1.0), Inches(3.7), Inches(11.3), Inches(2.9))
    tf_proto = tb_proto.text_frame
    tf_proto.word_wrap = True
    p = tf_proto.paragraphs[0]
    p.text = "严格的科研实证准则与指标口径"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_DARK
    p.space_after = Pt(8)

    proto_pts = [
        "双卡算力环境：双卡 NVIDIA RTX 4090 D (48GB VRAM × 2)，188GB 内存，保证多进程无显存溢出并行仿真。",
        "荣和对局率口径 (ron_game_rate)：P0 通过抓对手点炮获胜的对局占比，按整局统计，避免与单张弃牌胡率混淆。",
        "严格 Oracle FFR：仅统计对手向听数 ≤ 1 且 P0 自身确实未听牌（纯 Bluff）时，对手产生实际向听数倒退的极端事件。",
        "尾部风险 CVaR-10%：统计智能体在最差 10% 恶劣手牌和危险局势下的期望得分，衡量系统防崩盘能力。"
    ]
    for pt in proto_pts:
        p_pt = tf_proto.add_paragraph()
        p_pt.text = "• " + pt
        p_pt.font.size = Pt(12)
        p_pt.font.color.rgb = COLOR_TEXT_DARK
        p_pt.space_after = Pt(6)

    # ==========================================
    # SLIDE 10: 核心主实验结果 (E2)
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "主实验结果 (E2)：四级系统阶梯全面超越强基线", "9. MAIN EXPERIMENTAL RESULTS")

    t_shape = s10.shapes.add_table(5, 7, Inches(0.8), Inches(1.6), Inches(11.733), Inches(3.2))
    t = t_shape.table
    widths = [Inches(2.5), Inches(1.5), Inches(1.5), Inches(1.6), Inches(1.5), Inches(1.6), Inches(1.533)]
    for j, w in enumerate(widths):
        t.columns[j].width = w

    headers = ["系统方案", "平均得分 ↑", "胡牌率 ↑", "荣和对局率 ↑", "放铳率 (DR) ↓", "excess-DIR ↑", "CVaR-10% ↑"]
    for j, h in enumerate(headers):
        cell = t.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_BG_DARK
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEXT_LIGHT
        p.alignment = PP_ALIGN.CENTER

    data_e2 = [
        ("L0 (Qwen-base)", "13.479", "28.67%", "19.66%", "19.34%", "-6.620pp", "-87.995"),
        ("L1 (Reactive-z)", "14.352", "31.05%", "21.40%", "18.92%", "-4.875pp", "-84.540"),
        ("GateValue (No-deceive)", "19.328", "38.20%", "26.28%", "16.91%", "0.000", "-74.075"),
        ("MASK-full (Ours)", "24.893", "51.15%", "34.73%", "15.37%", "+8.450pp", "-58.390"),
    ]
    for i, row in enumerate(data_e2):
        for j, val in enumerate(row):
            cell = t.cell(i+1, j)
            cell.fill.solid()
            if i == 3:
                cell.fill.fore_color.rgb = RGBColor(230, 247, 255)
            else:
                cell.fill.fore_color.rgb = COLOR_CARD_BG if i % 2 == 0 else RGBColor(255, 255, 255)
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(12)
            p.alignment = PP_ALIGN.CENTER
            if i == 3:
                p.font.bold = True
                p.font.color.rgb = COLOR_PRIMARY if j == 0 else COLOR_GREEN
            else:
                p.font.color.rgb = COLOR_TEXT_DARK

    # Bottom Highlight Box
    add_card(s10, Inches(0.8), Inches(5.1), Inches(11.733), Inches(1.8), border_color=COLOR_PRIMARY)
    tb_hl = s10.shapes.add_textbox(Inches(1.0), Inches(5.2), Inches(11.3), Inches(1.6))
    tf_hl = tb_hl.text_frame
    tf_hl.word_wrap = True
    p = tf_hl.paragraphs[0]
    p.text = "关键配对显著性统计结论 (MASK-full vs No-deceive)"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(4)

    p_sub = tf_hl.add_paragraph()
    p_sub.text = (
        "• 局均期望得分提升 +5.565 分 (p = 3.62e-7)，显著击穿传统基线瓶颈！\n"
        "• 超额诱导荣和率 (excess-DIR) 暴增 +8.450pp (p = 2.99e-15)，证实动作塑形直接诱导对手点炮！\n"
        "• 自身放铳率 (DR) 降低 -1.545pp (p = 1.64e-7)，尾部抗灾得分 CVaR 改善 +15.685 分 (p = 7.01e-6)。"
    )
    p_sub.font.size = Pt(12)
    p_sub.font.color.rgb = COLOR_TEXT_DARK

    # ==========================================
    # SLIDE 11: 深度消融实验 (E3)
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "消融实验结果 (E3)：拆解各模块的独立必要性与贡献", "10. ABLATION STUDY")

    t_shape = s11.shapes.add_table(7, 6, Inches(0.8), Inches(1.6), Inches(11.733), Inches(5.2))
    t = t_shape.table
    widths = [Inches(2.5), Inches(1.4), Inches(1.4), Inches(1.6), Inches(1.6), Inches(3.233)]
    for j, w in enumerate(widths):
        t.columns[j].width = w

    headers = ["消融变体", "平均得分 ↑", "放铳率 ↓", "excess-DIR ↑", "CVaR-10% ↑", "学术判定与贡献归因"]
    for j, h in enumerate(headers):
        cell = t.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_BG_DARK
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEXT_LIGHT
        p.alignment = PP_ALIGN.CENTER

    data_e3 = [
        ("MASK-full (完整系统)", "24.893", "15.37%", "+8.450pp", "-58.390", "基准最优解，兼顾收益与尾部安全"),
        ("-deceive (禁用欺骗)", "19.328", "16.91%", "0.000", "-74.075", "得分骤降 5.56 分，证明欺骗是核心收益源"),
        ("-risk (常开欺骗/无门控)", "13.794", "17.95%", "+0.765pp", "-72.785", "得分大跌 11.1 分！放铳激增，证明门控至关重要"),
        ("-course (无课程微调)", "20.874", "16.04%", "+4.355pp", "-67.090", "规划能力下降，证明 LLM 课程学习不可或缺"),
        ("-B_phi (无公开信念)", "24.879", "15.70%", "+8.320pp", "-57.895", "得分微降，但 DR 恶化 +0.33pp (p<1e-3)，充当安全盾"),
        ("-z (无启发特征)", "22.045", "14.98%", "+10.050pp", "-58.580", "权衡点：得分减少 2.85 分，防守略保守")
    ]
    for i, row in enumerate(data_e3):
        for j, val in enumerate(row):
            cell = t.cell(i+1, j)
            cell.fill.solid()
            if i == 0:
                cell.fill.fore_color.rgb = RGBColor(230, 247, 255)
            elif i == 2:
                cell.fill.fore_color.rgb = RGBColor(255, 231, 230)
            else:
                cell.fill.fore_color.rgb = COLOR_CARD_BG if i % 2 == 0 else RGBColor(255, 255, 255)
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(11)
            p.font.name = "Microsoft YaHei"
            if j < 5:
                p.alignment = PP_ALIGN.CENTER
            if i == 0:
                p.font.bold = True

    # ==========================================
    # SLIDE 12: 经验风险前沿与反欺骗压力测试 (E5)
    # ==========================================
    s12 = prs.slides.add_slide(blank_layout)
    add_header(s12, "经验风险前沿与压力测试 (E5)：对抗环境下的鲁棒性边界", "11. RISK FRONTIER & STRESS TEST")

    # Left Card: Normal Holdout
    add_card(s12, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    tb_l = s12.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.2), Inches(4.8))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True
    p = tf_l.paragraphs[0]
    p.text = "常规冻结对手：门控阈值调优前沿"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(10)

    pts_e5_l = [
        "Default 设置最优：平均分 24.893，CVaR -58.390，Deceive 触发率 11.71%。",
        "Tight (紧门控)：防守极稳 (DR 15.57%)，但略微损失微小进攻机会 (得分 24.490)。",
        "Loose (宽门控)：得分 24.411，excess-DIR 达 +8.68pp，适合追求极致进攻。",
        "Constant (恒常欺骗)：彻底崩塌！得分仅 13.794，说明盲目欺骗在常规池不可行。"
    ]
    for pt in pts_e5_l:
        p_pt = tf_l.add_paragraph()
        p_pt.text = "• " + pt
        p_pt.font.size = Pt(12)
        p_pt.font.color.rgb = COLOR_TEXT_DARK
        p_pt.space_after = Pt(8)

    # Right Card: Anti-deception
    add_card(s12, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    tb_r = s12.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.3), Inches(4.8))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True
    p = tf_r.paragraphs[0]
    p.text = "Anti-Deception 强对抗压力测试"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = COLOR_RED
    p.space_after = Pt(10)

    pts_e5_r = [
        "对手具备防欺骗机制：针对 P0 常见虚假信号做惩罚反制，全场放铳率激增至 ~28%。",
        "No-deceive 无法生存：在强对抗环境下，不开欺骗平均得分跌入负值 (-0.326)。",
        "动态门控自适应生存：Default/Loose 依然维持正得分 (+2.60 ~ +2.77)，CVaR 显著改善 13+ 分！",
        "常开欺骗灾难性失分：Constant-deceive 暴跌至 -8.036 分，再次验证门控价值网络的护栏价值。"
    ]
    for pt in pts_e5_r:
        p_pt = tf_r.add_paragraph()
        p_pt.text = "✔ " + pt
        p_pt.font.size = Pt(12)
        p_pt.font.color.rgb = COLOR_TEXT_DARK
        p_pt.space_after = Pt(8)

    # ==========================================
    # SLIDE 13: 机理审计与神经对手验证 (E6 & V4)
    # ==========================================
    s13 = prs.slides.add_slide(blank_layout)
    add_header(s13, "机理审计与神经对手泛化：实测决策与外部有效性", "12. MECHANISM & GENERALIZATION")

    # Left: Belief drift
    add_card(s13, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    tb_l = s13.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.2), Inches(4.8))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True
    p = tf_l.paragraphs[0]
    p.text = "E6 机理审计：271,517 步实测决策"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(10)

    e6_pts = [
        "信念正向漂移验证：在 31,788 次触发 Deceive 的决策中，对手信念产生 +0.02103 的显著正向偏移。",
        "正向成功率对比：Deceive 诱导成功率高达 62.73%（对比常态 Exploit 的 42.03%），机理上实证了动作塑形有效扭曲了对手认知。",
        "单步推理延迟：中位数 p50 仅为 246.6 ms，p95 为 446.4 ms，具备准实时在线部署能力。"
    ]
    for pt in e6_pts:
        p_pt = tf_l.add_paragraph()
        p_pt.text = "• " + pt
        p_pt.font.size = Pt(12)
        p_pt.font.color.rgb = COLOR_TEXT_DARK
        p_pt.space_after = Pt(10)

    # Right: Neural Opponent V4
    add_card(s13, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    tb_r = s13.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.3), Inches(4.8))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True
    p = tf_r.paragraphs[0]
    p.text = "V4 独立风险感知神经对手验证"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = COLOR_GREEN
    p.space_after = Pt(10)

    v4_pts = [
        "脱离规则教师：对手基于深度学习独立感知 P0 听牌风险，拥有 Aggressive / Balanced / Cautious 风格。",
        "Balanced 对手表现：Full vs No-deceive 平均得分 +1.89 分，excess-DIR 达 +6.2pp (p=0.0107)，放铳率降 3.2pp。",
        "Cautious 对手表现：Full 相对 No-deceive excess-DIR 达 +6.1pp (p=0.0079)，放铳率显著下降 -4.1pp (p=0.0038)。",
        "外部有效性确立：证明 MASK 不依赖特定规则对手，面对现代深度学习对手依然保持强劲收益！"
    ]
    for pt in v4_pts:
        p_pt = tf_r.add_paragraph()
        p_pt.text = "✔ " + pt
        p_pt.font.size = Pt(12)
        p_pt.font.color.rgb = COLOR_TEXT_DARK
        p_pt.space_after = Pt(8)

    # ==========================================
    # SLIDE 14: 总结与学术启发
    # ==========================================
    s14 = prs.slides.add_slide(blank_layout)
    add_header(s14, "总结与学术启发：不完全信息博弈的认知进阶", "13. CONCLUSION & INSIGHTS")

    # 3 Bottom Summary Cards
    conclusions = [
        ("① 诱捕优于驱赶",
         "博弈欺骗的真正价值不在于吓退对手 (FFR)，而在于伪装安全诱导对手做出错误进攻并点炮 (excess-DIR)。盲目追求高 FFR 反而陷入流局惩罚。",
         COLOR_PRIMARY),
        ("② 时机重于频率",
         "常开欺骗必败，门控时机为王。只有当门控价值网络计算出显著的边际优势时激活欺骗，才能将'虚张声势'转化为'确定性期望收益'。",
         COLOR_GOLD),
        ("③ 大模型与轻量网络的优雅分工",
         "LLM 擅长全局常识与长期规划，轻量 MLP 擅长敏锐风控与实时数值仲裁。分层解耦架构为复杂系统大模型落地提供了通用范式。",
         COLOR_GREEN),
    ]
    for i, (title, content, col) in enumerate(conclusions):
        left = Inches(0.8 + i * 4.0)
        card = add_card(s14, left, Inches(1.6), Inches(3.73), Inches(5.2))
        
        bar = s14.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, Inches(1.6), Inches(3.73), Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = col
        bar.line.fill.background()

        tb = s14.shapes.add_textbox(left + Inches(0.2), Inches(1.9), Inches(3.33), Inches(4.7))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXT_DARK
        p.space_after = Pt(14)

        p2 = tf.add_paragraph()
        p2.text = content
        p2.font.size = Pt(13)
        p2.font.color.rgb = COLOR_TEXT_MUTED
        p2.space_after = Pt(10)

    # Save presentation
    output_path = "/home-students/yu0666/majiang/majiang_ai/mahjong-ai-battle-main/MASK_Paper_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

    # Also make a copy or link in the home root for quick user access
    home_path = "/home-students/yu0666/MASK_Paper_Presentation.pptx"
    try:
        import shutil
        shutil.copy2(output_path, home_path)
        print(f"Copied to home: {home_path}")
    except Exception as e:
        print(f"Could not copy to home: {e}")

if __name__ == "__main__":
    create_deck()
